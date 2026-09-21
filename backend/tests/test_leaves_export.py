"""Tests des exports congés (GET /api/leaves/export.xlsx et .docx) : accès
admin uniquement, contenu généré, filtre de période optionnel.

Isolation : même pattern que les autres tests API — importe app.py avec cwd
pointé sur un dossier temporaire.

Lancer : python -m unittest tests.test_leaves_export -v   (depuis backend/)
"""
import os
import sys
import unittest
from io import BytesIO

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # backend/
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))  # tests/ (for _app_loader)

os.environ.setdefault('SECRET_KEY', 'test-secret-key-for-unittests-only')

from _app_loader import load_fresh_app
ohmapp = load_fresh_app('ohmflow_leaves_export_test_')


class LeavesExportTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with ohmapp.app.app_context():
            admin = ohmapp.User.query.filter_by(username='Admin').first()
            admin.must_change_password = False
            admin.mfa_enabled = True
            worker = ohmapp.User(username='ExportWorker', role='user', must_change_password=False,
                                  pin_hash=ohmapp.generate_password_hash('unused'))
            ohmapp.db.session.add(worker)
            ohmapp.db.session.commit()
            cls.admin_id = admin.id
            cls.worker_id = worker.id
            cls.admin_token = ohmapp.serializer.dumps({'user_id': admin.id})
            cls.worker_token = ohmapp.serializer.dumps({'user_id': worker.id})

            ohmapp.db.session.add_all([
                ohmapp.Leave(user_id=worker.id, type='CONGE', date_start='2026-03-10', date_end='2026-03-12',
                             status='APPROVED', days_count=3),
                ohmapp.Leave(user_id=worker.id, type='MALADIE', date_start='2026-06-01', date_end='2026-06-01',
                             status='PENDING', days_count=1),
            ])
            ohmapp.db.session.commit()

        cls.admin_client = ohmapp.app.test_client()
        cls.admin_client.set_cookie(ohmapp.COOKIE_NAME, cls.admin_token)
        cls.worker_client = ohmapp.app.test_client()
        cls.worker_client.set_cookie(ohmapp.COOKIE_NAME, cls.worker_token)

    def test_xlsx_export_requires_admin(self):
        res = self.worker_client.get('/api/leaves/export.xlsx')
        self.assertEqual(res.status_code, 403)

    def test_docx_export_requires_admin(self):
        res = self.worker_client.get('/api/leaves/export.docx')
        self.assertEqual(res.status_code, 403)

    def test_xlsx_export_contains_all_leaves(self):
        res = self.admin_client.get('/api/leaves/export.xlsx')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(
            res.headers['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(res.data))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0], ('Employé', 'Type', 'Début', 'Fin', 'Jours', 'Statut'))
        self.assertEqual(len(rows), 3)  # header + 2 leaves
        body = {r[0:2] for r in rows[1:]}
        self.assertIn(('ExportWorker', 'Congé'), body)
        self.assertIn(('ExportWorker', 'Maladie'), body)

    def test_docx_export_contains_all_leaves(self):
        res = self.admin_client.get('/api/leaves/export.docx')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(
            res.headers['Content-Type'],
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )
        from docx import Document
        doc = Document(BytesIO(res.data))
        table = doc.tables[0]
        texts = [[c.text for c in row.cells] for row in table.rows]
        self.assertEqual(texts[0], ['Employé', 'Type', 'Début', 'Fin', 'Jours', 'Statut'])
        self.assertEqual(len(texts), 3)  # header + 2 leaves

    def test_period_filter_narrows_results(self):
        res = self.admin_client.get('/api/leaves/export.xlsx?date_start=2026-05-01&date_end=2026-12-31')
        self.assertEqual(res.status_code, 200)
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(res.data))
        rows = list(wb.active.iter_rows(values_only=True))
        self.assertEqual(len(rows), 2)  # header + only the June leave
        self.assertEqual(rows[1][1], 'Maladie')


if __name__ == '__main__':
    unittest.main()
